"""Build the portable demo data from data/rkg.sqlite.

Writes, next to this file:
  kg.sqlite                 slim, flattened copy of the graph for the Streamlit app (no full publication text)
  exports/BoE_MBF_KG.xlsx   the same tables as an Excel workbook, with a README sheet
  exports/csv/*.csv         the same tables as CSV (UTF-8 with BOM so Excel opens them cleanly)

Run from the project root:  .venv/bin/python demo/build_demo_data.py
The source database is opened read-only.
"""
import json
import re
import sqlite3
from collections import defaultdict
from pathlib import Path

import pandas as pd
import yaml
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "rkg.sqlite"
SEED = ROOT / "rkg" / "ontology" / "seed.yaml"
OUT = Path(__file__).resolve().parent
DB_OUT = OUT / "kg.sqlite"
XLSX = OUT / "exports" / "BoE_MBF_KG.xlsx"
CSV_DIR = OUT / "exports" / "csv"

SERIES = {"swp": "Staff Working Paper", "qb": "Quarterly Bulletin", "overground": "Bank Overground",
          "underground": "Bank Underground", "insight": "Bank Insight"}
ROLE = {1.0: "central", 0.6: "analysed", 0.3: "mentioned"}
MIN_EMERGENT_DOCS = 2  # same rule as rkg/graph_tools.py: model-proposed concepts need 2+ publications to be shown

src = sqlite3.connect(f"file:{SRC}?mode=ro", uri=True)
CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(df):
    """PDF text keeps the 'ff' ligature as \\x1b ('a\\x1bect' = 'affect'); Excel rejects control characters."""
    for col in df.columns:
        if not (pd.api.types.is_string_dtype(df[col]) or df[col].dtype == object):
            continue
        df[col] =df[col].map(lambda v: CTRL.sub("", v.replace("\x1b", "ff")) if isinstance(v, str) else v)
    return df


q = lambda sql, *a: clean(pd.read_sql_query(sql, src, params=a))  # noqa: E731

# ---------------------------------------------------------------- publications (the 245 in the graph)
docs = q("SELECT id doc_id, title, date, source, url, authors FROM documents "
         "WHERE kept=1 AND id IN (SELECT doc_id FROM extractions) ORDER BY date")
summ = q("SELECT src doc_id, props FROM edges WHERE rel='SUMMARY'")
summ["summary"] = summ.props.map(lambda p: json.loads(p or "{}").get("text", ""))
docs = docs.merge(summ[["doc_id", "summary"]], on="doc_id", how="left")
docs["year"] = docs.date.str[:4].astype(int)
docs["series"] = docs.source.map(SERIES)
docs["authors"] = docs.authors.map(lambda a: "; ".join(json.loads(a or "[]")))
doc_ids = set(docs.doc_id)

# ---------------------------------------------------------------- concepts
concepts = q("SELECT id concept_id, name, facet, definition, aliases, seeded FROM concepts")
concepts["aliases"] = concepts.aliases.map(lambda a: "; ".join(json.loads(a or "[]")))
dc = q("SELECT doc_id, concept_id, salience, quote FROM doc_concepts")
dc = dc[dc.doc_id.isin(doc_ids)]
n_any = dc.groupby("concept_id").doc_id.nunique()
concepts["shown"] = concepts.apply(lambda c: bool(c.seeded) or n_any.get(c.concept_id, 0) >= MIN_EMERGENT_DOCS, axis=1)
seed = {c["id"]: c for c in yaml.safe_load(SEED.read_text())["concepts"]}
concepts["broader"] = concepts.concept_id.map(lambda c: seed.get(c, {}).get("broader"))
concepts["event_date"] = concepts.concept_id.map(lambda c: seed.get(c, {}).get("date"))

hier = q("SELECT src broader, dst narrower FROM edges WHERE rel='BROADER_THAN'")
prereq = q("SELECT src learn_first, dst concept FROM edges WHERE rel='PREREQUISITE_OF'")

# ---------------------------------------------------------------- claims
claims = q("SELECT id claim_id, doc_id, date, driver, driver_raw, sign direction, outcome, outcome_raw, channel, "
           "conditions, evidence_type, hedging, quote, quote_how quote_check FROM mechanisms ORDER BY date, id")
claims["quote_check"] = claims.quote_check.replace("", "no quote found")

# Manual review: directions recorded the wrong way round by the extractor (the quote and channel say otherwise).
CORRECTIONS = {
    2983: {"driver_raw": "Increase in credit spreads", "direction": "increases"},  # quote: spread rise is amplified
    3275: {"direction": "increases"},  # quote: further price falls necessitate more sales
}
for cid, fix in CORRECTIONS.items():
    for col, val in fix.items():
        claims.loc[claims.claim_id == cid, col] = val
name = dict(zip(concepts.concept_id, concepts.name))

# ---------------------------------------------------------------- relations between claims
rel = q("SELECT src, rel, dst, props FROM edges WHERE rel IN ('CONTRADICTS','QUALIFIES','SUPPORTS')")
rel = pd.DataFrame({
    "later_claim_id": rel.src.str[5:].astype(int), "earlier_claim_id": rel.dst.str[5:].astype(int),
    "relation": rel.rel.map({"CONTRADICTS": "flagged as contradiction", "QUALIFIES": "depends on context",
                             "SUPPORTS": "same or supporting claim"}),
    "reason": rel.props.map(lambda p: json.loads(p or "{}").get("reason", "")),
})
# pairs touching a corrected claim: re-label by the corrected directions
direction = dict(zip(claims.claim_id, claims.direction))
for i, r in rel[rel.later_claim_id.isin(CORRECTIONS) | rel.earlier_claim_id.isin(CORRECTIONS)].iterrows():
    same = direction[r.later_claim_id] == direction[r.earlier_claim_id]
    if same and r.relation != "same or supporting claim":
        rel.at[i, "relation"] = "same or supporting claim"
        rel.at[i, "reason"] = ""  # as for other same-direction pairs, which the pipeline does not send for checking
    elif not same and r.relation == "same or supporting claim":
        print(f"check pair {r.later_claim_id}-{r.earlier_claim_id}: now opposite directions")

# ---------------------------------------------------------------- other links
events = concepts[concepts.facet == "event"][["concept_id", "name", "event_date", "definition"]] \
    .rename(columns={"concept_id": "event_id", "event_date": "date"})
pub_events = q("SELECT src doc_id, dst event_id FROM edges WHERE rel='RESPONDS_TO'")
pub_methods = q("SELECT src doc_id, dst method_id FROM edges WHERE rel='USES_METHOD'")
all_titles = q("SELECT id, title, date, source FROM documents")
cites = q("SELECT src citing_doc_id, dst cited_doc_id FROM edges WHERE rel='CITES' AND src != dst")
cites = cites.merge(all_titles.rename(columns={"id": "cited_doc_id", "title": "cited_title", "date": "cited_date",
                                               "source": "cited_series"}), on="cited_doc_id", how="left")
cites["cited_series"] = cites.cited_series.map(SERIES)
cites["cited_in_graph"] = cites.cited_doc_id.isin(doc_ids)
oq = q("SELECT id question_id, doc_id, date, text question, concepts FROM open_questions ORDER BY date")
oq["concepts"] = oq.concepts.map(lambda c: "; ".join(json.loads(c or "[]")))

# MBF share of each series by year (denominator: every publication found when the corpus was screened)
corpus = q("SELECT substr(date,1,4) year, source, COUNT(*) all_publications, "
           "SUM(kept=1 AND id IN (SELECT doc_id FROM extractions)) mbf_publications FROM documents "
           "WHERE date IS NOT NULL GROUP BY 1, 2")
corpus["series"] = corpus.source.map(SERIES)
corpus = corpus.drop(columns="source")

# ---------------------------------------------------------------- slim database for the app
if DB_OUT.exists():
    DB_OUT.unlink()
out = sqlite3.connect(DB_OUT)
tables = {
    "publications": docs[["doc_id", "title", "date", "year", "series", "url", "authors", "summary"]],
    "concepts": concepts, "concept_hierarchy": hier, "learning_path": prereq,
    "publication_concepts": dc, "claims": claims, "claim_relations": rel,
    "events": events, "publication_events": pub_events, "publication_methods": pub_methods,
    "citations": cites, "open_questions": oq, "corpus_by_year": corpus,
}
for t, df in tables.items():
    df.to_sql(t, out, index=False)
out.execute("CREATE INDEX ix_claims_link ON claims(driver, outcome)")
out.commit()
out.close()


# ---------------------------------------------------------------- readable tables for Excel / CSV
def family(cid):
    kids = defaultdict(list)
    for b, n in hier.itertuples(index=False):
        kids[b].append(n)
    out_, stack = {cid}, [cid]
    while stack:
        for k in kids[stack.pop()]:
            if k not in out_:
                out_.add(k)
                stack.append(k)
    return out_


title = dict(zip(docs.doc_id, docs.title))
series = dict(zip(docs.doc_id, docs.series))
url = dict(zip(docs.doc_id, docs.url))

x_claims = pd.DataFrame({
    "claim_id": claims.claim_id, "date": claims.date,
    "driver": claims.driver.map(lambda c: name.get(c, c)), "direction": claims.direction,
    "outcome": claims.outcome.map(lambda c: name.get(c, c)),
    "driver_as_written": claims.driver_raw, "outcome_as_written": claims.outcome_raw,
    "channel": claims.channel, "conditions": claims.conditions,
    "evidence_type": claims.evidence_type, "hedging": claims.hedging,
    "quote": claims.quote, "quote_check": claims.quote_check,
    "publication": claims.doc_id.map(title), "series": claims.doc_id.map(series), "url": claims.doc_id.map(url),
    "driver_id": claims.driver, "outcome_id": claims.outcome, "doc_id": claims.doc_id,
})

g = claims.groupby(["driver", "outcome"])
links = pd.DataFrame({
    "publications": g.doc_id.nunique(), "claims": g.size(),
    "first_year": g.date.min().str[:4], "last_year": g.date.max().str[:4],
    "increases": g.direction.apply(lambda s: (s == "increases").sum()),
    "decreases": g.direction.apply(lambda s: (s == "decreases").sum()),
    "no_effect_or_ambiguous": g.direction.apply(lambda s: s.isin(["no_effect", "ambiguous"]).sum()),
    "share_measured_in_data": g.evidence_type.apply(lambda s: round((s == "empirical").mean(), 2)),
    "share_stated_firmly": g.hedging.apply(lambda s: round((s == "firm").mean(), 2)),
}).reset_index()
links = links[links.driver != links.outcome]
links.insert(0, "outcome_name", links.outcome.map(lambda c: name.get(c, c)))
links.insert(0, "driver_name", links.driver.map(lambda c: name.get(c, c)))
links = links.rename(columns={"driver": "driver_id", "outcome": "outcome_id",
                              "driver_name": "driver", "outcome_name": "outcome"})
links = links.sort_values(["publications", "claims"], ascending=False)

n_central = dc[dc.salience >= 0.6].groupby("concept_id").doc_id.nunique()
n_claims = pd.concat([claims.driver, claims.outcome]).value_counts()
x_concepts = concepts.assign(
    publications_central_or_analysed=concepts.concept_id.map(n_central).fillna(0).astype(int),
    claims=concepts.concept_id.map(n_claims).fillna(0).astype(int),
    vocabulary=concepts.seeded.map({1: "fixed vocabulary (57)", 0: "proposed by the model"}),
    broader=concepts.broader.map(lambda c: name.get(c, c) if c else ""),
)[["concept_id", "name", "facet", "vocabulary", "definition", "aliases", "broader",
   "publications_central_or_analysed", "claims", "event_date", "shown"]] \
    .sort_values(["vocabulary", "publications_central_or_analysed"], ascending=[True, False])

x_pubs = docs.assign(
    claims=docs.doc_id.map(claims.doc_id.value_counts()).fillna(0).astype(int),
    events=docs.doc_id.map(pub_events.assign(n=pub_events.event_id.map(name))
                           .groupby("doc_id").n.agg("; ".join)).fillna(""),
    central_concepts=docs.doc_id.map(dc[dc.salience >= 1.0].assign(n=dc.concept_id.map(lambda c: name.get(c, c)))
                                     .groupby("doc_id").n.agg("; ".join)).fillna(""),
)[["doc_id", "title", "date", "series", "authors", "claims", "events", "central_concepts", "summary", "url"]]

x_pubconc = dc.assign(publication=dc.doc_id.map(title), date=dc.doc_id.map(dict(zip(docs.doc_id, docs.date))),
                      concept=dc.concept_id.map(lambda c: name.get(c, c)), role=dc.salience.map(ROLE))[
    ["date", "publication", "concept", "role", "quote", "doc_id", "concept_id"]].sort_values(["date", "publication"])


def claim_side(df, col, prefix):
    c = x_claims.set_index("claim_id")
    return pd.DataFrame({
        f"{prefix}_claim": df[col].map(lambda i: f"{c.at[i, 'driver']} {c.at[i, 'direction']} {c.at[i, 'outcome']}"),
        f"{prefix}_quote": df[col].map(c.quote), f"{prefix}_publication": df[col].map(c.publication),
        f"{prefix}_date": df[col].map(c.date),
    })


opp = rel[rel.relation != "same or supporting claim"].reset_index(drop=True)
x_opposing = pd.concat([opp[["relation"]], claim_side(opp, "later_claim_id", "later"),
                        claim_side(opp, "earlier_claim_id", "earlier"),
                        opp[["reason", "later_claim_id", "earlier_claim_id"]]], axis=1) \
    .sort_values(["relation", "later_date"], ascending=[False, False])

x_events = pub_events.assign(event=pub_events.event_id.map(name), publication=pub_events.doc_id.map(title),
                             date=pub_events.doc_id.map(dict(zip(docs.doc_id, docs.date))))[
    ["event", "date", "publication", "doc_id", "event_id"]].sort_values(["event", "date"])
x_cites = cites.assign(citing_title=cites.citing_doc_id.map(title))[
    ["citing_title", "cited_title", "cited_date", "cited_series", "cited_in_graph", "citing_doc_id", "cited_doc_id"]]
x_oq = oq.assign(publication=oq.doc_id.map(title),
                 concepts=oq.concepts.map(lambda s: "; ".join(name.get(c, c) for c in s.split("; ") if c)))[
    ["date", "question", "concepts", "publication", "doc_id"]].sort_values("date", ascending=False)

# coverage matrix: sector x vulnerability, publications where both are central or analysed (incl. narrower)
main = dc[dc.salience >= 0.6]
docs_of = main.groupby("concept_id").doc_id.agg(set)
shown = concepts[concepts.shown]


def fam_docs(c):
    return set().union(*(docs_of.get(x, set()) for x in family(c)))


sectors = shown[(shown.facet == "sector") & (shown.concept_id != "mbf")].concept_id
vulns = shown[shown.facet == "vulnerability"].concept_id
x_matrix = pd.DataFrame([[len(fam_docs(s) & fam_docs(v)) for v in vulns] for s in sectors],
                        index=[name[s] for s in sectors], columns=[name[v] for v in vulns])
x_matrix.index.name = "sector \\ vulnerability"

x_corpus = corpus.pivot_table(index="year", columns="series", values=["mbf_publications", "all_publications"],
                              aggfunc="sum", fill_value=0)
x_corpus.columns = [f"{a} - {b}" for a, b in x_corpus.columns]
x_corpus["mbf_publications - all series"] = corpus.groupby("year").mbf_publications.sum()
x_corpus["all_publications - all series"] = corpus.groupby("year").all_publications.sum()
x_corpus["mbf_share"] = (x_corpus["mbf_publications - all series"] / x_corpus["all_publications - all series"]).round(3)
x_corpus = x_corpus.reset_index()

sheets = {
    "Publications": x_pubs, "Claims": x_claims, "Causal links": links, "Concepts": x_concepts,
    "Publication concepts": x_pubconc, "Opposing claims": x_opposing, "Events": x_events,
    "Citations": x_cites, "Open questions": x_oq, "Coverage matrix": x_matrix.reset_index(),
    "Corpus by year": x_corpus,
}

n_seed = int(concepts.seeded.sum())
readme = [
    ("Bank of England market-based finance (MBF) research: knowledge graph", ""),
    ("", ""),
    ("What this is", "A knowledge graph built from 245 Bank of England staff publications on market-based finance "
     "since 2015 (Staff Working Papers, Quarterly Bulletin, Bank Overground, Bank Underground, Bank Insights). "
     "Each causal claim is recorded as driver → direction → outcome, with its evidence type, hedging, date and a "
     "supporting quote, and tied to its publication."),
    ("Built", "Claims were extracted by a language model (Claude Haiku 4.5) and quotes checked against the "
     "source text. Full publication text is not included; follow the url column to read a publication."),
    ("", ""),
    ("Sheet", "What it contains"),
    ("Publications", f"{len(x_pubs)} publications: title, date, series, authors, number of claims, events analysed, "
     "central concepts, a model-written summary and the link."),
    ("Claims", f"{len(x_claims):,} causal claims. driver/outcome are concepts from the vocabulary; "
     "'as written' columns give the publication's own wording. Filter on driver and outcome to see the evidence "
     "for a link."),
    ("Causal links", f"{len(links):,} driver → outcome links with counts of publications and claims, years, "
     "directions and the share of claims measured in data / stated firmly. Counts here are exact concept pairs; "
     "the report's Table 1 also includes narrower concepts (see Concepts → broader)."),
    ("Concepts", f"{len(x_concepts)} concepts: the fixed vocabulary of {n_seed} (with synonyms) plus concepts the "
     "model proposed. 'shown' = used by the app (fixed vocabulary, or proposed and found in 2+ publications)."),
    ("Publication concepts", "Which concepts each publication covers: central, analysed or mentioned, with a quote."),
    ("Opposing claims", f"{len(x_opposing)} pairs of claims pointing in opposite directions on the same link, "
     "each with the automated check's verdict and reason."),
    ("Events", "Publications analysing each stress episode (e.g. March 2020 dash for cash, 2022 LDI episode)."),
    ("Citations", "Links from publications to other Bank publications they cite."),
    ("Open questions", f"{len(x_oq)} open questions raised by authors, newest first."),
    ("Coverage matrix", "Publications in which both a sector (rows) and a vulnerability (columns) are central or "
     "analysed. Zeros are research gaps within this collection."),
    ("Corpus by year", "All publications screened and MBF publications kept, by year and series."),
    ("", ""),
    ("Known limitations", ""),
    ("1", "Direction, hedging, evidence type and concept matching were assigned by a language model and have not "
     "been reviewed by experts. Treat counts as a guide to where evidence sits, not as a quality score."),
    ("2", "The 'US Treasuries' concept also picks up generic statements about interest rates."),
    ("3", "About 14% of quotes are the closest source passage rather than the model's exact quote "
     "(quote_check = located); one claim has no quote."),
    ("4", "Working papers were read from their abstract, introduction and conclusion. Staff publications set out "
     "their authors' views, not those of the Bank's policy committees."),
]

XLSX.parent.mkdir(parents=True, exist_ok=True)
CSV_DIR.mkdir(parents=True, exist_ok=True)
for f in CSV_DIR.glob("*.csv"):
    f.unlink()

HEAD = PatternFill("solid", fgColor="EEF3FA")
LONG = {"summary", "quote", "reason", "channel", "conditions", "definition", "question", "aliases",
        "later_quote", "earlier_quote", "central_concepts"}
with pd.ExcelWriter(XLSX, engine="openpyxl") as xw:
    pd.DataFrame(readme, columns=["A", "B"]).to_excel(xw, sheet_name="README", index=False, header=False)
    ws = xw.sheets["README"]
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
        if row[0].value in ("What this is", "Built", "Sheet", "Known limitations"):
            row[0].font = Font(bold=True)
            if row[0].value == "Sheet":
                row[1].font = Font(bold=True)

    for sheet, df in sheets.items():
        df.to_excel(xw, sheet_name=sheet, index=False)
        ws = xw.sheets[sheet]
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for j, col in enumerate(df.columns, start=1):
            c = ws.cell(row=1, column=j)
            c.font, c.fill = Font(bold=True), HEAD
            c.alignment = Alignment(wrap_text=True, vertical="top")
            sample = df[col].astype(str).head(200)
            width = max(len(str(col)), int(sample.str.len().quantile(0.9)) if len(sample) else 10)
            long = col in LONG or col in ("title", "publication", "cited_title", "citing_title")
            ws.column_dimensions[get_column_letter(j)].width = min(max(width, 8) + 2, 70 if long else 32)
            if long:
                for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                    for x in cell:
                        x.alignment = Alignment(wrap_text=True, vertical="top")
        if "url" in df.columns:
            j = list(df.columns).index("url") + 1
            for r in range(2, len(df) + 2):
                cell = ws.cell(row=r, column=j)
                if cell.value:
                    cell.hyperlink, cell.style = cell.value, "Hyperlink"
        df.to_csv(CSV_DIR / f"{sheet.lower().replace(' ', '_')}.csv", index=False, encoding="utf-8-sig")

print(f"wrote {DB_OUT.relative_to(ROOT)} ({DB_OUT.stat().st_size / 1e6:.1f} MB)")
print(f"wrote {XLSX.relative_to(ROOT)} ({XLSX.stat().st_size / 1e6:.1f} MB)")
print(f"wrote {len(list(CSV_DIR.glob('*.csv')))} CSVs to {CSV_DIR.relative_to(ROOT)}")
for s, df in sheets.items():
    print(f"  {s:22s} {len(df):>6,} rows")
